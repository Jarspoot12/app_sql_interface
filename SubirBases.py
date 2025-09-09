# -*- coding: utf-8 -*- 
# -*- coding: utf-8 -*-
import pandas as pd
import os
from sqlalchemy import create_engine, inspect, text
from datetime import datetime, date
import openpyxl
import hashlib
import logging
import re
from typing import Dict, List, Tuple, Optional 

"""ESTE CODIGO TRANSFORMA ARCHIVOS EXCEL CON DIFERENTES ESTRUCTURAS HISTÓRICAS (2015-2024) 
A UN FORMATO UNIFICADO Y LUEGO LOS DIVIDE EN 3 TABLAS RELACIONADAS PARA POSTGRESQL"""

# Configuración de logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('procesamiento_split.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Configuración de la conexión a PostgreSQL local (mi pc) | pc servidor del c5
DB_USER = 'app_ri_user' # postgres
DB_PASSWORD = '1234'  # ComplejoC5
DB_HOST = 'localhost' # localhost del c5
DB_PORT = '5432' # 5432
DB_NAME = 'app_sql' # Complejo


# Crear la cadena de conexión
connection_string = f'postgresql+psycopg2://{DB_USER}:{DB_PASSWORD}@{DB_HOST}:{DB_PORT}/{DB_NAME}'
engine = create_engine(
    connection_string,
    connect_args={"options": "-c search_path=app_sql,public"},
    pool_pre_ping=True
)
# Mapeo de columnas por versión (igual que en el archivo original)
COLUMN_MAPPING = {
    # Estructura PRINCIPAL (destino en PostgreSQL)
    "principal": {
        "FOLIO": "FOLIO",
        "FECHA": "FECHA",
        "TELEFONO": "TELEFONO",
        "UBICACION": "UBICACION",
        "COLONIA": "COLONIA",
        "MUNICIPIO": "MUNICIPIO",
        "RCBD": "RCBD",
        "DESP": "DESP",
        "LLEG": "LLEG",
        "LIBR": "LIBR",
        "T1": "T1",
        "T2": "T2",
        "T3": "T3",
        "T4": "T4",
        "CORPORACION": "CORPORACION",
        "TIPO": "TIPO",
        "MAKEDESC": "MAKEDESC",
        "MODEL": "MODEL",
        "COLOR": "COLOR",
        "VYR": "VYR",
        "VLIC": "VLIC",
        "ST": "ST",
        "ADDITIONAL": "ADDITIONAL",
        "CLSDESC": "CLSDESC",
        "OPERADOR": "OPERADOR",
        "DESPACHADOR": "DESPACHADOR",
        "UNIDAD": "UNIDAD",
        "DIV": "DIV",
        "COMENTARIOS": "COMENTARIOS",
        "CHLNAME": "CHLNAME",
        "CHFNAME": "CHFNAME",
        "ORIGEN": "ORIGEN",
        "LATITUD": "LATITUD",
        "LONGITUD": "LONGITUD",
        "PROCEDENTE": "PROCEDENTE",
        "MTVOCIERRE": "MTVOCIERRE",
        "NOTACIERRE": "NOTACIERRE",
        "SECTOR": "SECTOR",
        "NOTASUSR": "NOTASUSR",
        "PERSONASINV": "PERSONASINV",
        "VEHICULOSINV": "VEHICULOSINV",
        "TMPTIPIFICACION": "TMPTIPIFICACION",
        "TMPDESPACHO": "TMPDESPACHO",
    },
    
    # Mapeo para archivos 2024 inicios porque despues de marzo ya sigue como la version principal:
    "2024": {
        "FOLIO": ("FOLIO_LLAMADA", "A"),
        "FECHA": ("FECHA_LLAMADA", "B"),
        "TELEFONO": ("NUMERO_TELEFONO", "AB"),
        "UBICACION": ("REFERENCIAS", "M"),
        "COLONIA": ("COLONIA", "N"),
        "MUNICIPIO": ("MUNICIPIO", "L"),
        "RCBD": ("HORA_LLAMADA", "I"),
        "DESP": ("TIEMPO_DESPACHO", "AS"),
        "LLEG": ("TIEMPO_LLEGADA", "AT"),
        "LIBR": ("TIEMPO_SOLUCION", "AU"),
        "CORPORACION": ("NOMBRE_CORPORACION", "AN"),
        "TIPO": ("TIPO DE INCIDENTE", "X"),
        "MAKEDESC": ("MARCA", "BQ"),
        "MODEL": ("MODELO", "BR"),
        "COLOR": ("COLOR", "BV"),
        "VYR": ("ANIO", "BU"),
        "VLIC": ("PLACA", "BN"),
        "ADDITIONAL": ("TIPO_VEHICULO", "BP"),
        "CLSDESC": ("RAZONAMIENTO_CORPORACION", "AE"),
        "OPERADOR": ("TELEFONISTA", "AY"),
        "DESPACHADOR": ("NOMBRE_RADIO_OPERADOR", "AZ"),
        "UNIDAD": ("UNIDAD", "CO"),
        "COMENTARIOS": ("DESCRIPCION_DE_LA_LLAMADA + RESPONSABLE_DE_UNIDAD", ["AM", "AP"]),
        "CHFNAME": ("NOMBRE_DENUNCIANTE", "BG"),
        "ORIGEN": ("ORIGEN_LLAMADA", "AA"),
        "LATITUD": ("COORDENADA_X", "S"),
        "LONGITUD": ("COORDENADA_Y", "T")
    },
    
    # Mapeo para archivos del 2015 al 2023
    "2015-2023": {
        "FOLIO": ("FOLIO_LLAMADA", "S"),
        "FECHA": ("FECHA_LLAMADA", "B"),
        "TELEFONO": ("NUMERO_TELEFONO", "D"),
        "UBICACION": ("REFERENCIAS", "J"),
        "COLONIA": ("COLONIA", "K"),
        "MUNICIPIO": ("MUNICIPIO", "I"),
        "RCBD": ("HORA_LLAMADA", "C"),
        "DESP": ("TIEMPO_DESPACHO", "AC"),
        "LLEG": ("TIEMPO_LLEGADA", "AD"),
        "LIBR": ("TIEMPO_SOLUCION", "AE"),
        "CORPORACION": ("NOMBRE_CORPORACION", "R"),
        "TIPO": ("TIPO DE INCIDENTE", "E"),
        "CLSDESC": ("RAZONAMIENTO_CORPORACION", "X"),
        "OPERADOR": ("TELEFONISTA", "AI"),
        "DESPACHADOR": ("NOMBRE_RADIO_OPERADOR", "AJ"),
        "COMENTARIOS": ("DESCRIPCION_DE_LA_LLAMADA + RESPONSABLE_DE_UNIDAD", ["Q", "Y"]),
        "CHFNAME": ("NOMBRE_DENUNCIANTE", "AT"),
        "ORIGEN": ("ORIGEN_LLAMADA", "AO"),
        "LATITUD": ("COORDENADA_X", "O"),
        "LONGITUD": ("COORDENADA_Y", "P")
    }
}

def calculate_file_hash(file_path):
    """Calcula el hash MD5 de un archivo"""
    hash_md5 = hashlib.md5()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()

def get_processed_files():
    """Obtiene la lista de archivos ya procesados desde la base de datos"""
    try:
        with engine.connect() as conn:
            # Crear tabla de control si no existe
            create_table_query = text("""
                CREATE TABLE IF NOT EXISTS processed_files_split (
                    id SERIAL PRIMARY KEY,
                    filename VARCHAR(255),
                    file_hash VARCHAR(32),
                    processed_date TIMESTAMP,
                    version_estructura VARCHAR(50),
                    filas_principales INTEGER,
                    filas_corporaciones INTEGER,
                    filas_comentarios INTEGER
                )
            """)
            conn.execute(create_table_query)
            conn.commit()
            
            # Obtener lista de archivos procesados
            select_query = text("SELECT filename, file_hash FROM processed_files_split")
            result = conn.execute(select_query)
            return {row[0]: row[1] for row in result}
    except Exception as e:
        logger.error(f"Error al obtener archivos procesados: {str(e)}")
        return {}

def is_excel_file(filename):
    """Verifica si el archivo es un Excel válido (ignora archivos temporales)"""
    return (filename.endswith(('.xlsx', '.xls')) and 
            not filename.startswith('~$'))

def detect_version_structure(filename: str, headers: List[str]) -> str:
    """
    Detecta la versión de estructura basándose en el nombre del archivo y los headers
    """
    # Extraer año del nombre del archivo
    year_match = re.search(r'20(1[5-9]|2[0-4])', filename)
    if year_match:
        year = int(year_match.group())
        if year == 2024:
            # Detección más específica para 2024
            estructura_2024_indicadores = [
                "FOLIO_LLAMADA", "NUMERO_TELEFONO", "COORDENADA_X", "COORDENADA_Y",
                "FECHA_LLAMADA", "HORA_LLAMADA", "TIEMPO_DESPACHO", "TIEMPO_LLEGADA",
                "TIEMPO_SOLUCION", "NOMBRE_CORPORACION", "TIPO DE INCIDENTE"
            ]
            
            # Contar cuántos indicadores de estructura 2024 están presentes
            indicadores_encontrados = sum(1 for header in headers if any(indicator in header for indicator in estructura_2024_indicadores))
            
            # Si encuentra al menos 5 indicadores, es estructura 2024 temprana
            if indicadores_encontrados >= 5:
                logger.info(f"   - Indicadores 2024 encontrados: {indicadores_encontrados}")
                return "2024"
            else:
                logger.info(f"   - Indicadores 2024 encontrados: {indicadores_encontrados} (insuficientes)")
                return "principal"
                
        elif 2015 <= year <= 2023:
            return "2015-2023"
    
    # Si no se puede detectar por nombre, intentar por headers
    # Detección por estructura de headers
    estructura_2024_indicadores = [
        "FOLIO_LLAMADA", "NUMERO_TELEFONO", "COORDENADA_X", "COORDENADA_Y",
        "FECHA_LLAMADA", "HORA_LLAMADA", "TIEMPO_DESPACHO", "TIEMPO_LLEGADA",
        "TIEMPO_SOLUCION", "NOMBRE_CORPORACION", "TIPO DE INCIDENTE"
    ]
    
    estructura_principal_indicadores = [
        "FOLIO", "TELEFONO", "LATITUD", "LONGITUD", "FECHA", "CORPORACION"
    ]
    
    # Contar indicadores de cada estructura
    indicadores_2024 = sum(1 for header in headers if any(indicator in header for indicator in estructura_2024_indicadores))
    indicadores_principal = sum(1 for header in headers if any(indicator in header for indicator in estructura_principal_indicadores))
    
    logger.info(f"   - Indicadores 2024: {indicadores_2024}")
    logger.info(f"   - Indicadores principal: {indicadores_principal}")
    
    # Determinar por mayoría
    if indicadores_2024 > indicadores_principal:
        return "2024"
    elif indicadores_principal > indicadores_2024:
        return "principal"
    else:
        # Si no hay diferencia clara, usar 2015-2023 como default
        return "2015-2023"

def get_column_index(column_ref: str) -> int:
    """Convierte referencia de columna Excel (A, B, C...) a índice numérico"""
    result = 0
    for char in column_ref:
        result = result * 26 + (ord(char.upper()) - ord('A') + 1)
    return result - 1

def extract_column_value(row_data: List, column_ref: str) -> str:
    """Extrae valor de una columna específica usando referencia Excel"""
    try:
        index = get_column_index(column_ref)
        if index < len(row_data):
            value = row_data[index]
            if value is None:
                return ""
            elif isinstance(value, (datetime, date)):
                return str(value)
            else:
                return str(value)
        return ""
    except:
        return ""

def combine_columns(row_data: List, column_refs: List[str]) -> str:
    """Combina múltiples columnas en un solo valor"""
    values = []
    for ref in column_refs:
        value = extract_column_value(row_data, ref)
        if value and value.strip():
            values.append(value.strip())
    return " | ".join(values)

def transform_row_data(row_data: List, version: str, headers: List[str] = None) -> Dict:
    """
    Transforma una fila de datos según la versión de estructura
    """
    transformed = {}
    mapping = COLUMN_MAPPING.get(version, {})
    
    for target_col, source_info in mapping.items():
        if isinstance(source_info, str):
            # Mapeo directo (estructura principal) - buscar en headers
            if headers:
                try:
                    # Buscar el índice de la columna en los headers
                    col_index = headers.index(source_info)
                    if col_index < len(row_data):
                        value = row_data[col_index]
                        # Convertir a string de forma segura
                        if value is None:
                            transformed[target_col] = ""
                        elif isinstance(value, (datetime, date)):
                            transformed[target_col] = str(value)
                        else:
                            transformed[target_col] = str(value)
                    else:
                        transformed[target_col] = ""
                except ValueError:
                    # Si no encuentra la columna, usar valor vacío
                    transformed[target_col] = ""
            else:
                transformed[target_col] = ""
        elif isinstance(source_info, tuple):
            if len(source_info) == 2:
                if isinstance(source_info[1], str):
                    # Mapeo simple con referencia de columna
                    value = extract_column_value(row_data, source_info[1])
                    transformed[target_col] = str(value) if value else ""
                elif isinstance(source_info[1], list):
                    # Mapeo de columnas combinadas
                    value = combine_columns(row_data, source_info[1])
                    transformed[target_col] = str(value) if value else ""
            else:
                transformed[target_col] = ""
        else:
            transformed[target_col] = ""
    
    # Agregar campos de metadatos
    transformed['fecha_carga'] = datetime.now()
    transformed['version_estructura'] = str(version)
    
    return transformed

def create_split_tables():
    """Crea las 3 tablas separadas con la estructura optimizada y relaciones"""
    try:
        with engine.connect() as conn:
            # 1. Tabla PRINCIPAL (folios únicos) - TABLA PADRE
            create_principal_query = text("""
                CREATE TABLE IF NOT EXISTS principal (
                    id SERIAL PRIMARY KEY,
                    FOLIO TEXT UNIQUE NOT NULL,
                    FECHA DATE,
                    TELEFONO TEXT,
                    UBICACION TEXT,
                    COLONIA TEXT,
                    MUNICIPIO TEXT,
                    TIPO TEXT,
                    MAKEDESC TEXT,
                    MODEL TEXT,
                    COLOR TEXT,
                    VYR TEXT,
                    VLIC TEXT,
                    ST TEXT,
                    ADDITIONAL TEXT,
                    CLSDESC TEXT,
                    OPERADOR TEXT,
                    DESPACHADOR TEXT,
                    UNIDAD TEXT,
                    DIV TEXT,
                    CHLNAME TEXT,
                    CHFNAME TEXT,
                    ORIGEN TEXT,
                    LATITUD TEXT,
                    LONGITUD TEXT,
                    PROCEDENTE TEXT,
                    SECTOR TEXT,
                    PERSONASINV TEXT,
                    VEHICULOSINV TEXT,
                    fecha_carga TIMESTAMP,
                    version_estructura TEXT,
                    origen_archivo TEXT
                )
            """)
            conn.execute(create_principal_query)
            
            # 2. Tabla CORPORACIONES (múltiples corporaciones por folio) - TABLA HIJA
            create_corporaciones_query = text("""
                CREATE TABLE IF NOT EXISTS corporaciones (
                    id SERIAL PRIMARY KEY,
                    FOLIO TEXT NOT NULL,
                    CORPORACION TEXT,
                    RCBD TEXT,
                    DESP TEXT,
                    LLEG TEXT,
                    LIBR TEXT,
                    T1 TEXT,
                    T2 TEXT,
                    T3 TEXT,
                    T4 TEXT,
                    TMPTIPIFICACION TEXT,
                    TMPDESPACHO TEXT,
                    fecha_carga TIMESTAMP,
                    CONSTRAINT fk_corporaciones_principal 
                    FOREIGN KEY (FOLIO) REFERENCES principal(FOLIO) 
                    ON DELETE CASCADE ON UPDATE CASCADE
                )
            """)
            conn.execute(create_corporaciones_query)
            
            # 3. Tabla COMENTARIOS (comentarios por folio único) - TABLA HIJA
            create_comentarios_query = text("""
                CREATE TABLE IF NOT EXISTS comentarios (
                    id SERIAL PRIMARY KEY,
                    FOLIO TEXT UNIQUE NOT NULL,
                    COMENTARIOS TEXT,
                    MTVOCIERRE TEXT,
                    NOTACIERRE TEXT,
                    NOTASUSR TEXT,
                    fecha_carga TIMESTAMP,
                    CONSTRAINT fk_comentarios_principal
                    FOREIGN KEY (FOLIO) REFERENCES principal(FOLIO)
                    ON DELETE CASCADE ON UPDATE CASCADE
                )
            """)
            conn.execute(create_comentarios_query)
            
            # Crear índices para optimizar JOINs y consultas
            try:
                # Índices para la tabla principal
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_principal_folio ON principal(FOLIO)"))
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_principal_fecha ON principal(FECHA)"))
                
                # Índices para la tabla corporaciones
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_corporaciones_folio ON corporaciones(FOLIO)"))
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_corporaciones_fecha ON corporaciones(fecha_carga)"))
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_corporaciones_corporacion ON corporaciones(CORPORACION)"))
                
                # Índices para la tabla comentarios
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_comentarios_folio ON comentarios(FOLIO)"))
                conn.execute(text("CREATE INDEX IF NOT EXISTS idx_comentarios_fecha ON comentarios(fecha_carga)"))
                
                logger.info("OK: Índices creados/verificados exitosamente")
            except Exception as e:
                logger.warning(f"WARNING: Error creando índices (pueden ya existir): {str(e)}")
            
            conn.commit()
            logger.info("OK: Tablas separadas con relaciones creadas/verificadas exitosamente")
            logger.info("   - PRINCIPAL: Tabla padre con FOLIO único")
            logger.info("   - CORPORACIONES: Tabla hija con FK a PRINCIPAL")
            logger.info("   - COMENTARIOS: Tabla hija con FK a PRINCIPAL")
            logger.info("   - CASCADE: Eliminación/actualización automática en tablas hijas")
    except Exception as e:
        logger.error(f"ERROR: Error creando tablas separadas: {str(e)}")
        raise

def split_data_into_tables(df: pd.DataFrame, filename: str) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Divide los datos unificados en las 3 tablas separadas
    """
    logger.info(f"SPLIT: Dividiendo datos de {filename} en 3 tablas...")
    
    # 1. Tabla PRINCIPAL - Obtener filas únicas por FOLIO (manteniendo la primera aparición)
    # Agrupar por FOLIO y tomar la primera fila de cada grupo
    df_principal = df.drop_duplicates(subset=['folio'], keep='first').copy()
    
    # Seleccionar solo las columnas que van a la tabla principal
    columnas_principal = [
        'folio','fecha','telefono','ubicacion','colonia','municipio',
        'tipo','makedesc','model','color','vyr','vlic','st','additional','clsdesc',
        'operador','despachador','unidad','div','chlname','chfname','origen',
        'latitud','longitud','procedente','sector','personasinv','vehiculosinv',
        'fecha_carga','version_estructura','origen_archivo'
    ]
    
    # Filtrar solo las columnas que existen en el DataFrame
    columnas_existentes_principal = [col for col in columnas_principal if col in df_principal.columns]
    df_principal = df_principal[columnas_existentes_principal]
    
    # 2. Tabla CORPORACIONES - Obtener todas las filas con sus tiempos por corporación
    columnas_corporaciones = [
        'folio','corporacion','rcbd','desp','lleg','libr',
        't1','t2','t3','t4','tmptipificacion','tmpdespacho','fecha_carga'
    ]
    
    # Filtrar solo las columnas que existen en el DataFrame
    columnas_existentes_corp = [col for col in columnas_corporaciones if col in df.columns]
    df_corporaciones = df[columnas_existentes_corp].copy()
    
    # 3. Tabla COMENTARIOS - Obtener comentarios únicos por folio
    # Agrupar por FOLIO y combinar comentarios únicos
    df_comentarios = df.groupby('folio').agg({
        'comentarios': lambda x: ' | '.join(filter(None, set(x.astype(str)))),
        'mtvocierre': lambda x: ' | '.join(filter(None, set(x.astype(str)))) if 'mtvocierre' in df.columns else '',
        'notacierre': lambda x: ' | '.join(filter(None, set(x.astype(str)))) if 'notacierre' in df.columns else '',
        'notasusr': lambda x: ' | '.join(filter(None, set(x.astype(str)))) if 'notasusr' in df.columns else ''
    }).reset_index()
    
    # Agregar fecha_carga
    df_comentarios['fecha_carga'] = datetime.now()
    
    # Agregar logging para ver qué columnas faltan
    columnas_faltantes_principal = set(columnas_principal) - set(columnas_existentes_principal)
    columnas_faltantes_corp = set(columnas_corporaciones) - set(columnas_existentes_corp)
    
    if columnas_faltantes_principal:
        logger.info(f"   - Columnas faltantes en tabla principal: {columnas_faltantes_principal}")
    if columnas_faltantes_corp:
        logger.info(f"   - Columnas faltantes en tabla corporaciones: {columnas_faltantes_corp}")
    
    logger.info(f"   - Filas principales (únicas): {len(df_principal)}")
    logger.info(f"   - Filas corporaciones: {len(df_corporaciones)}")
    logger.info(f"   - Filas comentarios: {len(df_comentarios)}")
    
    return df_principal, df_corporaciones, df_comentarios

def get_existing_folios():
    """Obtiene los folios que ya existen en las tablas para evitar duplicados"""
    try:
        with engine.connect() as conn:
            # Obtener folios existentes en tabla principal
            result_principal = conn.execute(text("SELECT DISTINCT FOLIO FROM principal"))
            folios_principal = {row[0] for row in result_principal}
            
            # Obtener folios existentes en tabla comentarios
            result_comentarios = conn.execute(text("SELECT DISTINCT FOLIO FROM comentarios"))
            folios_comentarios = {row[0] for row in result_comentarios}
            
            # Obtener folios existentes en tabla corporaciones
            result_corporaciones = conn.execute(text("SELECT DISTINCT FOLIO FROM corporaciones"))
            folios_corporaciones = {row[0] for row in result_corporaciones}
            
            return folios_principal, folios_comentarios, folios_corporaciones
    except Exception as e:
        logger.warning(f"WARNING: Error obteniendo folios existentes: {str(e)}")
        return set(), set(), set()

def filter_new_data(df_principal: pd.DataFrame, df_corporaciones: pd.DataFrame, df_comentarios: pd.DataFrame, 
                   existing_folios_principal: set, existing_folios_comentarios: set, existing_folios_corporaciones: set) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Filtra los datos para incluir solo folios nuevos que no existen en las tablas
    """
    # Filtrar tabla principal - solo folios nuevos
    df_principal_new = df_principal[~df_principal['folio'].isin(existing_folios_principal)].copy()
    
    # Filtrar tabla comentarios - solo folios nuevos
    df_comentarios_new = df_comentarios[~df_comentarios['folio'].isin(existing_folios_comentarios)].copy()
    
    # Para corporaciones, incluir todos los registros de folios nuevos
    folios_nuevos = set(df_principal_new['folio'])
    df_corporaciones_new = df_corporaciones[df_corporaciones['folio'].isin(folios_nuevos)].copy()
    
    logger.info(f"   - Filas principales nuevas: {len(df_principal_new)}")
    logger.info(f"   - Filas corporaciones nuevas: {len(df_corporaciones_new)}")
    logger.info(f"   - Filas comentarios nuevos: {len(df_comentarios_new)}")
    
    return df_principal_new, df_corporaciones_new, df_comentarios_new

def process_excel_file_split(file_path: str) -> bool:
    """
    Procesa un archivo Excel y lo transforma al formato unificado, luego lo divide en 3 tablas
    """
    try:
        filename = os.path.basename(file_path)
        logger.info(f"\nPROCESANDO: Procesando archivo: {filename}")

        # Cargar archivo Excel
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        
        # Obtener headers
        headers = []
        for cell in next(sheet.rows):
            if cell.value is not None:
                headers.append(str(cell.value).strip())
            else:
                headers.append(f'columna_{len(headers) + 1}')

        # Detectar versión de estructura
        version = detect_version_structure(filename, headers)
        logger.info(f"VERSION: Versión detectada: {version}")

        # Transformar datos
        transformed_data = []
        row_count = 0
        
        for row in sheet.rows:
            if row_count == 0:  # Saltar header
                row_count += 1
                continue
                
            row_data = [cell.value for cell in row]
            if any(x is not None for x in row_data):
                transformed_row = transform_row_data(row_data, version, headers)
                transformed_row['origen_archivo'] = filename
                transformed_data.append(transformed_row)
            row_count += 1

        if not transformed_data:
            logger.warning(f"WARNING: No se encontraron datos válidos en {filename}")
            return False

        # Crear DataFrame unificado
        df_unified = pd.DataFrame(transformed_data)
        df_unified.columns = [c.strip().lower() for c in df_unified.columns] # para mantener las columnas en minúsculas
        df_unified['fecha'] = pd.to_datetime(df_unified['fecha'], errors='coerce').dt.date


        # Convertir todas las columnas a string (excepto fecha_carga que es timestamp)
        for col in df_unified.columns:
            if col != 'fecha_carga':
                df_unified[col] = df_unified[col].astype(str)
        
        # Dividir datos en las 3 tablas
        df_principal, df_corporaciones, df_comentarios = split_data_into_tables(df_unified, filename)
        
        # Obtener folios existentes para evitar duplicados
        existing_folios_principal, existing_folios_comentarios, existing_folios_corporaciones = get_existing_folios()
        
        # Filtrar solo datos nuevos
        df_principal_new, df_corporaciones_new, df_comentarios_new = filter_new_data(
            df_principal, df_corporaciones, df_comentarios,
            existing_folios_principal, existing_folios_comentarios, existing_folios_corporaciones
        )
        
        # Solo cargar si hay datos nuevos
        if len(df_principal_new) == 0 and len(df_corporaciones_new) == 0 and len(df_comentarios_new) == 0:
            logger.info(f"INFO: No hay datos nuevos en {filename} - todos los folios ya existen")
            return True
        
        # Cargar datos a las 3 tablas (respetando dependencias de FK)
        try:
            # IMPORTANTE: Insertar en orden para respetar las claves foráneas
            # 1. PRIMERO: Tabla PRINCIPAL (tabla padre)
            if len(df_principal_new) > 0:
                df_principal_new.to_sql(
                    name='principal',
                    con=engine,
                    if_exists='append',
                    index=False,
                    chunksize=1000,
                    method=None
                )
                logger.info(f"OK: Tabla PRINCIPAL actualizada con {len(df_principal_new)} filas nuevas")
            
            # 2. SEGUNDO: Tabla CORPORACIONES (depende de PRINCIPAL)
            if len(df_corporaciones_new) > 0:
                df_corporaciones_new.to_sql(
                    name='corporaciones',
                    con=engine,
                    if_exists='append',
                    index=False,
                    chunksize=1000,
                    method=None
                )
                logger.info(f"OK: Tabla CORPORACIONES actualizada con {len(df_corporaciones_new)} filas nuevas")
            
            # 3. TERCERO: Tabla COMENTARIOS (depende de PRINCIPAL)
            if len(df_comentarios_new) > 0:
                df_comentarios_new.to_sql(
                    name='comentarios',
                    con=engine,
                    if_exists='append',
                    index=False,
                    chunksize=1000,
                    method=None
                )
                logger.info(f"OK: Tabla COMENTARIOS actualizada con {len(df_comentarios_new)} filas nuevas")
            
        except Exception as e:
            logger.warning(f"WARNING: Error con pandas to_sql, usando inserción directa con manejo de FK: {str(e)}")
            
            # Fallback: inserción directa con SQL respetando dependencias
            with engine.connect() as conn:
                try:
                    # 1. PRIMERO: Insertar en PRINCIPAL (tabla padre)
                    if len(df_principal_new) > 0:
                        columns_principal = list(df_principal_new.columns)
                        placeholders_principal = ', '.join([f':{col}' for col in columns_principal])
                        insert_principal_query = text(f"""
                            INSERT INTO principal ({', '.join(columns_principal)})
                            VALUES ({placeholders_principal})
                            ON CONFLICT (FOLIO) DO NOTHING
                        """)
                        
                        for _, row in df_principal_new.iterrows():
                            try:
                                row_dict = row.to_dict()
                                conn.execute(insert_principal_query, row_dict)
                            except Exception as insert_error:
                                logger.warning(f"WARNING: Error insertando en PRINCIPAL folio {row_dict.get('folio', 'N/A')}: {str(insert_error)}")
                                continue
                    
                    # 2. SEGUNDO: Insertar en CORPORACIONES (depende de PRINCIPAL)
                    if len(df_corporaciones_new) > 0:
                        columns_corporaciones = list(df_corporaciones_new.columns)
                        placeholders_corporaciones = ', '.join([f':{col}' for col in columns_corporaciones])
                        insert_corporaciones_query = text(f"""
                            INSERT INTO corporaciones ({', '.join(columns_corporaciones)})
                            VALUES ({placeholders_corporaciones})
                        """)
                        
                        for _, row in df_corporaciones_new.iterrows():
                            try:
                                row_dict = row.to_dict()
                                conn.execute(insert_corporaciones_query, row_dict)
                            except Exception as insert_error:
                                logger.warning(f"WARNING: Error insertando en CORPORACIONES folio {row_dict.get('folio', 'N/A')}: {str(insert_error)}")
                                continue
                    
                    # 3. TERCERO: Insertar en COMENTARIOS (depende de PRINCIPAL)
                    if len(df_comentarios_new) > 0:
                        columns_comentarios = list(df_comentarios_new.columns)
                        placeholders_comentarios = ', '.join([f':{col}' for col in columns_comentarios])
                        insert_comentarios_query = text(f"""
                            INSERT INTO comentarios ({', '.join(columns_comentarios)})
                            VALUES ({placeholders_comentarios})
                            ON CONFLICT (FOLIO) DO UPDATE SET
                            COMENTARIOS = EXCLUDED.COMENTARIOS,
                            fecha_carga = EXCLUDED.fecha_carga
                        """)
                        
                        for _, row in df_comentarios_new.iterrows():
                            try:
                                row_dict = row.to_dict()
                                conn.execute(insert_comentarios_query, row_dict)
                            except Exception as insert_error:
                                logger.warning(f"WARNING: Error insertando en COMENTARIOS folio {row_dict.get('folio', 'N/A')}: {str(insert_error)}")
                                continue
                    
                    conn.commit()
                    logger.info("OK: Inserción directa completada respetando dependencias FK")
                    
                except Exception as transaction_error:
                    conn.rollback()
                    logger.error(f"ERROR: Error en transacción, haciendo rollback: {str(transaction_error)}")
                    raise

        # Registrar archivo como procesado
        file_hash = calculate_file_hash(file_path)
        with engine.connect() as conn:
            conn.execute(
                text("""
                    INSERT INTO processed_files_split 
                    (filename, file_hash, processed_date, version_estructura, 
                     filas_principales, filas_corporaciones, filas_comentarios) 
                    VALUES (:filename, :file_hash, :processed_date, :version_estructura,
                           :filas_principales, :filas_corporaciones, :filas_comentarios)
                """),
                {
                    "filename": filename,
                    "file_hash": file_hash,
                    "processed_date": datetime.now(),
                    "version_estructura": version,
                    "filas_principales": len(df_principal_new),
                    "filas_corporaciones": len(df_corporaciones_new),
                    "filas_comentarios": len(df_comentarios_new)
                }
            )
            conn.commit()

        logger.info(f"OK: Archivo {filename} procesado y acumulado exitosamente")
        logger.info(f"   - Versión: {version}")
        logger.info(f"   - Filas unificadas originales: {len(transformed_data)}")
        logger.info(f"   - Filas principales nuevas: {len(df_principal_new)}")
        logger.info(f"   - Filas corporaciones nuevas: {len(df_corporaciones_new)}")
        logger.info(f"   - Filas comentarios nuevos: {len(df_comentarios_new)}")
        
        return True

    except Exception as e:
        logger.error(f"ERROR: Error procesando {filename}: {str(e)}")
        import traceback
        logger.error(traceback.format_exc())
        return False

def verify_integrity():
    """Verifica la integridad referencial de las tablas"""
    try:
        with engine.connect() as conn:
            # Verificar folios huérfanos en CORPORACIONES
            orphan_corporaciones = conn.execute(text("""
                SELECT COUNT(*) FROM corporaciones c 
                LEFT JOIN principal p ON c.FOLIO = p.FOLIO 
                WHERE p.FOLIO IS NULL
            """)).scalar()
            
            # Verificar folios huérfanos en COMENTARIOS
            orphan_comentarios = conn.execute(text("""
                SELECT COUNT(*) FROM comentarios c 
                LEFT JOIN principal p ON c.FOLIO = p.FOLIO 
                WHERE p.FOLIO IS NULL
            """)).scalar()
            
            # Obtener estadísticas generales
            total_principal = conn.execute(text("SELECT COUNT(*) FROM principal")).scalar()
            total_corporaciones = conn.execute(text("SELECT COUNT(*) FROM corporaciones")).scalar()
            total_comentarios = conn.execute(text("SELECT COUNT(*) FROM comentarios")).scalar()
            
            logger.info("INTEGRIDAD: Verificación de relaciones:")
            logger.info(f"   - Tabla PRINCIPAL: {total_principal} registros")
            logger.info(f"   - Tabla CORPORACIONES: {total_corporaciones} registros")
            logger.info(f"   - Tabla COMENTARIOS: {total_comentarios} registros")
            logger.info(f"   - Folios huérfanos en CORPORACIONES: {orphan_corporaciones}")
            logger.info(f"   - Folios huérfanos en COMENTARIOS: {orphan_comentarios}")
            
            if orphan_corporaciones == 0 and orphan_comentarios == 0:
                logger.info("✅ INTEGRIDAD: Todas las relaciones están correctas")
                return True
            else:
                logger.warning("⚠️ INTEGRIDAD: Se encontraron folios huérfanos")
                return False
                
    except Exception as e:
        logger.error(f"ERROR: Error verificando integridad: {str(e)}")
        return False

def main():
    """Función principal del proceso"""
    data_folder = 'DATA'
    logger.info("INICIANDO: Iniciando proceso de transformación, unificación y acumulación de datos...")

    # Verificar carpeta de datos
    if not os.path.exists(data_folder):
        logger.error(f"ERROR: La carpeta {data_folder} no existe")
        return

    # Crear tablas separadas (solo una vez) con relaciones
    create_split_tables()

    # Obtener archivos Excel
    excel_files = [f for f in os.listdir(data_folder) if is_excel_file(f)]
    if not excel_files:
        logger.warning(f"ERROR: No se encontraron archivos Excel válidos en la carpeta {data_folder}")
        return

    # Obtener archivos ya procesados
    processed_files = get_processed_files()
    logger.info(f"ARCHIVOS: Encontrados {len(excel_files)} archivos Excel para procesar")
    logger.info(f"PROCESADOS: Archivos ya procesados: {len(processed_files)}")

    # Procesar archivos
    processed = 0
    skipped = 0
    failed = 0

    for excel_file in excel_files:
        file_path = os.path.join(data_folder, excel_file)
        file_hash = calculate_file_hash(file_path)
        
        if excel_file in processed_files and processed_files[excel_file] == file_hash:
            logger.info(f"SALTANDO: Saltando archivo: {excel_file} (ya procesado)")
            skipped += 1
            continue
            
        if process_excel_file_split(file_path):
            processed += 1
        else:
            failed += 1

    # Verificar integridad de las relaciones
    verify_integrity()

    # Resumen final
    logger.info("\nRESUMEN: Resumen del proceso de acumulación:")
    logger.info(f"   - Archivos procesados: {processed}")
    logger.info(f"   - Archivos saltados: {skipped}")
    logger.info(f"   - Archivos con error: {failed}")
    logger.info(f"   - Total de archivos: {len(excel_files)}")
    logger.info("\nESTRUCTURA FINAL CON RELACIONES:")
    logger.info("   - Tabla PRINCIPAL: Folios únicos (TABLA PADRE)")
    logger.info("   - Tabla CORPORACIONES: Múltiples corporaciones por folio (FK → PRINCIPAL)")
    logger.info("   - Tabla COMENTARIOS: Comentarios únicos por folio (FK → PRINCIPAL)")
    logger.info("\nCARACTERÍSTICAS DE INTEGRIDAD:")
    logger.info("   - Claves foráneas con CASCADE (eliminación/actualización automática)")
    logger.info("   - Índices optimizados para JOINs eficientes")
    logger.info("   - Inserción ordenada respetando dependencias")
    logger.info("   - Manejo de conflictos y errores de integridad")
    logger.info("\nCONSULTAS OPTIMIZADAS:")
    logger.info("   - JOINs automáticos entre tablas")
    logger.info("   - Consultas con integridad referencial garantizada")
    logger.info("   - Eliminación automática de registros huérfanos")

if __name__ == "__main__":
    main()
    logger.info("\nCOMPLETADO: Proceso de acumulación con relaciones completado")